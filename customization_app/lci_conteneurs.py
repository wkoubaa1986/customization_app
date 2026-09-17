"""
Répartition d'une « Liste Commande Import » en conteneurs.

Un import se paie au conteneur, pas au mètre cube : savoir qu'une liste pèse
191 m³ ne dit pas si elle tient en trois 40' HC ni ce qu'il faut sortir pour y
arriver. Ce module range les lignes dans des conteneurs successifs, dans
l'ordre de la liste — donc famille par famille, puisque la liste est triée par
groupe — et SCINDE une ligne quand elle déborde : 500 raccords deviennent 300
dans le premier conteneur et 200 dans le second, ce qui évite de casser un
conteneur à moitié vide pour une seule référence.

La capacité retenue n'est pas le volume géométrique du conteneur : un 40' HC
cube 76 m³ mais on ne charge jamais que ~90 % (calage, palettes, formes
irrégulières). Le taux est réglable, la valeur par défaut est prudente.
"""

import json

import frappe
from frappe import _
from frappe.utils import cint, flt, formatdate

from customization_app.lci_observation import prix_retenu, qty_retenue
from customization_app.liste_commande_import import (
    DOCTYPE,
    _chat_json,
    _chunks,
    _guard,
    _strip_html,
    _target_rows,
)

# volume géométrique, en m³
TYPES = {
    "20'": 33.0,
    "40'": 67.0,
    "40' HC": 76.0,
}
TAUX_DEFAUT = 0.90       # part réellement chargeable
TYPE_DEFAUT = "40' HC"
MAX_CONTENEURS = 40      # garde-fou : une liste ne planifie pas une flotte


def capacite(type_conteneur=TYPE_DEFAUT, taux=TAUX_DEFAUT):
    """Volume chargeable d'un conteneur, arrondi au litre."""
    brut = TYPES.get(type_conteneur or TYPE_DEFAUT, TYPES[TYPE_DEFAUT])
    taux = flt(taux) or TAUX_DEFAUT
    return round(brut * taux, 3)


def repartir(lignes, cap, caps=None):
    """Range les lignes dans des conteneurs successifs.

    `lignes` : [{"row", "qty", "volume_unitaire", "montant_unitaire", "libelle"}]
    `cap`    : capacité par défaut, celle des conteneurs qu'on ouvre en cours de route
    `caps`   : capacités imposées, conteneur par conteneur — c'est ce qui permet
               de MÉLANGER les gabarits (un 20' derrière deux 40' HC, par exemple)
               sans obliger toute la commande à voyager dans le même format.

    Rend {"conteneurs": [...], "sans_volume": [...], "capacite": cap}.

    Le remplissage suit l'ordre reçu ; une ligne qui déborde est coupée à la
    quantité entière qui tient encore. Si même une unité ne tient pas, le
    conteneur est clos et la ligne recommence dans le suivant.
    """
    cap = flt(cap)
    if cap <= 0:
        frappe.throw(_("Capacité de conteneur invalide."))
    caps = [flt(c) for c in (caps or [])]

    conteneurs = []
    sans_volume = []

    def ouvrir():
        i = len(conteneurs)
        conteneurs.append({"no": i + 1,
                           "capacite": caps[i] if i < len(caps) and caps[i] > 0 else cap,
                           "volume": 0.0, "montant": 0.0, "lignes": []})
        return conteneurs[-1]

    courant = ouvrir()
    for l in lignes:
        qty = flt(l.get("qty"))
        vu = flt(l.get("volume_unitaire"))
        if qty <= 0:
            continue
        if vu <= 0:
            # volume inconnu : la ligne doit voyager, mais elle ne peut pas
            # entrer dans le calcul — on la signale au lieu de la deviner.
            sans_volume.append({"row": l.get("row"), "libelle": l.get("libelle"),
                                "qty": qty})
            continue

        reste = qty
        while reste > 0:
            if len(conteneurs) > MAX_CONTENEURS:
                frappe.throw(_("Plus de {0} conteneurs : vérifiez les volumes unitaires.")
                             .format(MAX_CONTENEURS))
            dispo = flt(courant["capacite"]) - courant["volume"]
            tient = int(dispo / vu) if vu > 0 else 0
            if tient <= 0:
                if not courant["lignes"]:
                    # rien ne tient dans un conteneur vide : une seule unité
                    # dépasse la capacité. On la charge quand même, seule.
                    tient = 1
                else:
                    courant = ouvrir()
                    continue
            part = min(reste, tient)
            courant["lignes"].append({
                "row": l.get("row"),
                "libelle": l.get("libelle"),
                "qty": part,
                "volume": round(part * vu, 4),
                "montant": round(part * flt(l.get("montant_unitaire")), 4),
                "scinde": part < qty,
            })
            courant["volume"] = round(courant["volume"] + part * vu, 4)
            courant["montant"] = round(courant["montant"] + part * flt(l.get("montant_unitaire")), 4)
            reste -= part

    if not conteneurs[-1]["lignes"]:
        conteneurs.pop()
    return {"conteneurs": conteneurs, "sans_volume": sans_volume, "capacite": cap}


def lignes_a_charger(doc):
    """Lignes réellement embarquées : les abandons ne voyagent pas."""
    out = []
    for r in doc.articles:
        if (r.decision or "") == "Abandonné":
            continue
        q = qty_retenue(r)
        if q <= 0:
            continue
        vol = flt(r.volume_ligne_m3)
        out.append({
            "row": r.name,
            "idx": r.idx,
            "libelle": (r.item_code or r.item_name or "")[:60],
            "item_name": r.item_name or "",
            "qty": q,
            "uom": r.uom or "",
            "volume_unitaire": vol / q if q else 0.0,   # additionnels compris
            "montant_unitaire": prix_retenu(r),
        })
    return out


@frappe.whitelist()
def plan_auto(docname, type_conteneur=TYPE_DEFAUT, taux=TAUX_DEFAUT, types=None,
              dates=None):
    """Calcule une répartition. N'écrit rien.

    `types` : gabarit imposé à chaque conteneur (liste JSON, C1 en premier).
    Ce qui dépasse cette liste part sur `type_conteneur`, le gabarit par défaut.
    `dates` : date de départ de chaque conteneur, même ordre. Un import part
    rarement d'un coup : C1 en 40' début octobre, C2 en 20' trois semaines plus
    tard. La date suit le conteneur partout — écran, onglet et Excel.
    """
    _guard()
    doc = frappe.get_doc(DOCTYPE, docname)
    if isinstance(types, str):
        types = json.loads(types or "[]")
    if isinstance(dates, str):
        dates = json.loads(dates or "[]")
    types = [t for t in (types or [])]
    dates = [d for d in (dates or [])]
    defaut = type_conteneur or TYPE_DEFAUT
    cap = capacite(defaut, taux)
    plan = repartir(lignes_a_charger(doc), cap,
                    [capacite(t or defaut, taux) for t in types])
    for i, c in enumerate(plan["conteneurs"]):
        c["type"] = (types[i] if i < len(types) and types[i] else defaut)
        c["date"] = (dates[i] if i < len(dates) else "") or ""
    plan["type"] = defaut
    plan["taux"] = flt(taux) or TAUX_DEFAUT
    plan["devise"] = doc.devise or "USD"
    return plan


@frappe.whitelist()
def appliquer_plan(docname, plan):
    """Écrit la répartition sur les lignes (`repartition_conteneurs`)."""
    _guard()
    if isinstance(plan, str):
        plan = json.loads(plan)
    doc = frappe.get_doc(DOCTYPE, docname)
    par_nom = {r.name: r for r in doc.articles}

    par_ligne = {}
    for c in plan.get("conteneurs") or []:
        for l in c.get("lignes") or []:
            par_ligne.setdefault(l.get("row"), []).append(
                {"no": c.get("no"), "qty": flt(l.get("qty"))})

    for row in doc.articles:
        parts = par_ligne.get(row.name)
        row.repartition_conteneurs = json.dumps(parts, ensure_ascii=False) if parts else ""

    doc.plan_conteneurs = json.dumps({
        "type": plan.get("type") or TYPE_DEFAUT,
        "taux": flt(plan.get("taux")) or TAUX_DEFAUT,
        "capacite": flt(plan.get("capacite")) or capacite(),
        "conteneurs": [{"no": c.get("no"),
                        "type": c.get("type") or plan.get("type") or TYPE_DEFAUT,
                        "date": c.get("date") or "",
                        "capacite": flt(c.get("capacite")) or flt(plan.get("capacite")),
                        "volume": flt(c.get("volume")),
                        "montant": flt(c.get("montant")),
                        "nb_lignes": len(c.get("lignes") or [])}
                       for c in plan.get("conteneurs") or []],
    }, ensure_ascii=False)
    doc.save(ignore_permissions=True)
    frappe.db.commit()
    return {"nb_conteneurs": len(plan.get("conteneurs") or []),
            "sans_volume": len(plan.get("sans_volume") or []),
            "gabarits": resume_gabarits(plan.get("conteneurs") or [])}


def resume_gabarits(conteneurs):
    """« 2 × 40' HC + 1 × 20' » — la flotte en une ligne."""
    ordre, compte = [], {}
    for c in conteneurs:
        t = c.get("type") or TYPE_DEFAUT
        if t not in compte:
            ordre.append(t)
        compte[t] = compte.get(t, 0) + 1
    return " + ".join(f"{compte[t]} × {t}" for t in ordre)


# ------------------------------------------------- relecture et feuille Excel

LABELS_CONT = {
    "Français": {"onglet": "Conteneurs", "titre": "Plan de chargement",
                 "conteneur": "Conteneur", "gabarit": "Gabarit", "capacite": "Capacité utile",
                 "remplissage": "Remplissage", "code": "Code", "designation": "Désignation",
                 "qty": "Qté", "uom": "UDM", "volume": "Volume (m³)", "montant": "Montant",
                 "total": "TOTAL", "scindee": "ligne scindée", "flotte": "Flotte",
                 "depart": "Départ",
                "sans_volume": "Sans volume unitaire — non réparties"},
    "English": {"onglet": "Containers", "titre": "Loading plan",
                "conteneur": "Container", "gabarit": "Type", "capacite": "Usable capacity",
                "remplissage": "Fill", "code": "Code", "designation": "Description",
                "qty": "Qty", "uom": "UoM", "volume": "Volume (m³)", "montant": "Amount",
                "total": "TOTAL", "scindee": "split line", "flotte": "Fleet",
                "depart": "Departure",
                "sans_volume": "No unit volume — not loaded"},
    "Deutsch": {"onglet": "Container", "titre": "Ladeplan",
                "conteneur": "Container", "gabarit": "Typ", "capacite": "Nutzbares Volumen",
                "remplissage": "Auslastung", "code": "Code", "designation": "Bezeichnung",
                "qty": "Menge", "uom": "ME", "volume": "Volumen (m³)", "montant": "Betrag",
                "total": "GESAMT", "scindee": "geteilte Position", "flotte": "Flotte",
                "depart": "Abfahrt",
                "sans_volume": "Ohne Stückvolumen — nicht verladen"},
    "Español": {"onglet": "Contenedores", "titre": "Plan de carga",
                "conteneur": "Contenedor", "gabarit": "Tipo", "capacite": "Capacidad útil",
                "remplissage": "Llenado", "code": "Código", "designation": "Descripción",
                "qty": "Cant.", "uom": "UdM", "volume": "Volumen (m³)", "montant": "Importe",
                "total": "TOTAL", "scindee": "línea dividida", "flotte": "Flota",
                "depart": "Salida",
                "sans_volume": "Sin volumen unitario — no cargadas"},
    "العربية": {"onglet": "الحاويات", "titre": "خطة الشحن",
                "conteneur": "حاوية", "gabarit": "النوع", "capacite": "السعة المفيدة",
                "remplissage": "نسبة الملء", "code": "المرجع", "designation": "التسمية",
                "qty": "الكمية", "uom": "الوحدة", "volume": "الحجم (م³)", "montant": "المبلغ",
                "total": "المجموع", "scindee": "بند مقسّم", "flotte": "الأسطول",
                "depart": "المغادرة",
                "sans_volume": "بدون حجم للوحدة — غير محمّلة"},
}


def plan_enregistre(doc):
    """Relit le plan appliqué : entête (`plan_conteneurs`) + parts des lignes.

    Rendu identique à `repartir` pour que l'affichage, l'onglet du formulaire et
    la feuille Excel montrent tous la même chose sans recalculer — un plan
    retouché à la main ne doit pas se « re-optimiser » tout seul à l'export.
    """
    try:
        entete = json.loads(doc.plan_conteneurs or "{}") or {}
    except Exception:
        entete = {}
    conteneurs = []
    par_no = {}
    for c in entete.get("conteneurs") or []:
        bloc = {"no": cint(c.get("no")), "type": c.get("type") or entete.get("type") or TYPE_DEFAUT,
                "date": c.get("date") or "",
                "capacite": flt(c.get("capacite")) or flt(entete.get("capacite")) or capacite(),
                "volume": 0.0, "montant": 0.0, "lignes": []}
        conteneurs.append(bloc)
        par_no[bloc["no"]] = bloc

    for r in doc.articles:
        try:
            parts = json.loads(r.repartition_conteneurs or "[]") or []
        except Exception:
            parts = []
        if not parts:
            continue
        q_tot = qty_retenue(r)
        vu = flt(r.volume_ligne_m3) / q_tot if q_tot else 0.0
        pu = prix_retenu(r)
        for part in parts:
            bloc = par_no.get(cint(part.get("no")))
            if not bloc:
                continue
            q = flt(part.get("qty"))
            bloc["lignes"].append({
                "row": r.name, "item_code": r.item_code or "",
                "libelle": r.item_name_traduit or r.item_name or r.item_code or "",
                "qty": q, "uom": r.uom or "",
                "volume": q * vu, "montant": q * pu,
                "scinde": len(parts) > 1,
            })
            bloc["volume"] += q * vu
            bloc["montant"] += q * pu

    return {"conteneurs": conteneurs, "type": entete.get("type") or TYPE_DEFAUT,
            "taux": flt(entete.get("taux")) or TAUX_DEFAUT,
            "capacite": flt(entete.get("capacite")) or capacite()}


def ecrire_feuille(wb, doc, plan=None):
    """Ajoute au classeur `wb` l'onglet du plan de chargement. Rend la feuille.

    Appelé par les DEUX exports (notre cotation et le fichier du fournisseur
    renvoyé annoté) : le fournisseur charge le conteneur, il lui faut la même
    feuille que nous.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    # unité de mesure dans la langue cible, comme le reste des exports
    # (import local : liste_commande_import est déjà importé par ce module)
    from customization_app.liste_commande_import import _labels as _labels_export
    from customization_app.liste_commande_import import _uom_out

    plan = plan or plan_enregistre(doc)
    if not plan.get("conteneurs"):
        return None
    L = LABELS_CONT.get(doc.langue_cible) or LABELS_CONT["Français"]
    LX = _labels_export(doc)
    devise = doc.devise or "USD"

    titre = L["onglet"][:31]
    if titre in wb.sheetnames:            # un renvoi suivant ne doit pas empiler les onglets
        del wb[titre]
    ws = wb.create_sheet(title=titre)
    if doc.langue_cible in ("العربية",):
        ws.sheet_view.rightToLeft = True

    for col, w in enumerate([6, 18, 52, 10, 8, 14, 14], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    gras = Font(bold=True)
    fond_ct = PatternFill("solid", fgColor="F9F0FF")
    fond_th = PatternFill("solid", fgColor="F0F2F5")

    vol_total = sum(flt(c["volume"]) for c in plan["conteneurs"])
    mnt_total = sum(flt(c["montant"]) for c in plan["conteneurs"])

    ws.append([f'{L["titre"]} — {doc.titre or doc.name} ({doc.name})'])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f'{L["flotte"]} : {resume_gabarits(plan["conteneurs"])}',
               "", "", "", "",
               round(vol_total, 3), round(mnt_total, 2)])
    ws.cell(row=2, column=6).font = gras
    ws.cell(row=2, column=7).font = gras
    ws.cell(row=2, column=7).number_format = f'#,##0.00 "{devise}"'
    ws.append([])

    for c in plan["conteneurs"]:
        cap = flt(c["capacite"])
        pct = (flt(c["volume"]) / cap * 100) if cap else 0
        ws.append([f'C{c["no"]}', c.get("type") or "",
                   (f'{L["depart"]} : {formatdate(c["date"])}' if c.get("date") else ""),
                   f'{L["capacite"]} : {cap:.2f} m³',
                   f'{L["remplissage"]} : {pct:.0f} %', round(flt(c["volume"]), 3),
                   round(flt(c["montant"]), 2)])
        ligne_ct = ws.max_row
        for col in range(1, 8):
            ws.cell(row=ligne_ct, column=col).font = gras
            ws.cell(row=ligne_ct, column=col).fill = fond_ct
        ws.cell(row=ligne_ct, column=7).number_format = f'#,##0.00 "{devise}"'

        ws.append(["#", L["code"], L["designation"], L["qty"], L["uom"],
                   L["volume"], L["montant"]])
        for col in range(1, 8):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.font = Font(bold=True, size=9)
            cell.fill = fond_th

        for i, l in enumerate(c["lignes"], 1):
            ws.append([i, l["item_code"],
                       l["libelle"] + (f' ({L["scindee"]})' if l["scinde"] else ""),
                       flt(l["qty"]), _uom_out(l["uom"], LX), round(flt(l["volume"]), 4),
                       round(flt(l["montant"]), 2)])
            ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True, vertical="center")
            ws.cell(row=ws.max_row, column=7).number_format = f'#,##0.00 "{devise}"'
        ws.append(["", "", L["total"], "", "", round(flt(c["volume"]), 3),
                   round(flt(c["montant"]), 2)])
        for col in (3, 6, 7):
            ws.cell(row=ws.max_row, column=col).font = gras
        ws.cell(row=ws.max_row, column=7).number_format = f'#,##0.00 "{devise}"'
        ws.append([])

    # ce qui ne voyage pas : dit ici, sinon l'absence passe pour un oubli
    sans = [r for r in doc.articles
            if (r.decision or "") != "Abandonné" and qty_retenue(r) > 0
            and not (r.repartition_conteneurs or "")]
    if sans:
        ws.append([L["sans_volume"]])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="FFB00020")
        for r in sans:
            ws.append(["", r.item_code or "",
                       r.item_name_traduit or r.item_name or "",
                       qty_retenue(r), _uom_out(r.uom, LX)])
    return ws


# --------------------------------------------- volumes manquants (estimation)

SYS_VOLUME = (
    "Tu estimes le VOLUME D'EXPÉDITION d'articles de traitement de l'eau "
    "(osmoseurs, filtres, membranes, raccords, vannes, pompes), emballage "
    "carton compris, pour calculer un chargement de conteneur.\n"
    "On te donne des REPÈRES : des articles voisins avec leur volume réel. "
    "Sers-t'en pour rester à l'échelle — un raccord en laiton de 1/4\" occupe "
    "quelques dizaines de cm³, pas 0,01 m³.\n"
    "Pour chaque article, rends le volume unitaire en m³, et si tu peux le "
    "conditionnement (pièces par carton et volume du carton en m³). Donne une "
    "base courte (la comparaison ou le calcul) et une confiance entre 0 et 1.\n"
    'Réponds en JSON : {"lignes": [{"i": <indice>, "volume_m3": <nombre>, '
    '"pcs_carton": <nombre ou null>, "volume_carton_m3": <nombre ou null>, '
    '"base": "...", "confiance": <0-1>}]}'
)

MAX_VOLUME_M3 = 5.0        # au-delà, ce n'est plus un article, c'est une erreur
LOT_VOLUME = 12            # articles par appel IA


def normaliser_estimation(volume_m3, pcs_carton=0, volume_carton_m3=0):
    """Ce qu'on garde d'une estimation : (volume, pcs/carton, CBM/carton) ou None.

    Trois garde-fous, parce qu'un volume inventé qui passe sans contrôle fausse
    un plan de chargement entier :
      - le volume peut se déduire du carton quand il manque ;
      - au-delà de MAX_VOLUME_M3 l'article n'existe pas, on ne propose rien ;
      - un conditionnement qui ne retombe pas sur le volume unitaire (±25 %)
        est abandonné : on garde le volume, on jette le carton.
    """
    vol = flt(volume_m3)
    pcs = flt(pcs_carton)
    vct = flt(volume_carton_m3)
    if vol <= 0 and pcs > 0 and vct > 0:
        vol = vct / pcs
    if vol <= 0 or vol > MAX_VOLUME_M3:
        return None
    if not (pcs > 0 and vct > 0 and abs(vct / pcs - vol) <= vol * 0.25):
        pcs, vct = 0.0, 0.0
    return vol, pcs, vct


def _reperes(doc, groupes):
    """Articles comparables DÉJÀ mesurés, groupe par groupe.

    C'est ce qui ancre l'estimation : sans repère, un modèle situe un raccord
    et un osmoseur dans le même ordre de grandeur.
    """
    out = {}
    for r in doc.articles:
        g = r.item_group or ""
        if g not in groupes or flt(r.volume_unitaire_m3) <= 0:
            continue
        out.setdefault(g, [])
        if len(out[g]) < 6:
            out[g].append({"article": r.item_name or r.item_code,
                           "volume_m3": round(flt(r.volume_unitaire_m3), 5)})
    return out


@frappe.whitelist()
def ai_estimer_volumes(docname, row_names=None):
    """Propose un volume pour les lignes qui n'en ont pas. N'ÉCRIT RIEN.

    Deux sources, dans l'ordre : la fiche Article quand elle porte déjà un
    volume (exact, gratuit), puis l'IA pour le reste. L'utilisateur tranche
    ligne par ligne dans l'aperçu — un volume inventé qui entre sans contrôle
    fausse un plan de chargement entier.
    """
    _guard()
    doc = frappe.get_doc(DOCTYPE, docname)
    cibles = [r for r in _target_rows(doc, row_names)
              if flt(r.volume_unitaire_m3) <= 0
              and (r.decision or "") != "Abandonné"
              and qty_retenue(r) > 0]
    if not cibles:
        return {"lignes": [], "message": _("Toutes les lignes ont déjà un volume.")}

    propositions = []
    restants = []
    codes = [r.item_code for r in cibles if r.item_code]
    connus = {}
    if codes:
        for it in frappe.get_all("Item", filters={"name": ("in", codes)},
                                 fields=["name", "custom_volume_m3"]):
            if flt(it.custom_volume_m3) > 0:
                connus[it.name] = flt(it.custom_volume_m3)

    for r in cibles:
        vol = connus.get(r.item_code or "")
        if vol:
            propositions.append({
                "row": r.name, "idx": r.idx, "item_code": r.item_code or "",
                "item_name": r.item_name or "", "qty": qty_retenue(r),
                "uom": r.uom or "", "volume_m3": vol,
                "pcs_carton": None, "volume_carton_m3": None,
                "base": _("fiche article (volume déjà connu)"),
                "confiance": 1, "estime": 0,
            })
        else:
            restants.append(r)

    reperes = _reperes(doc, {r.item_group or "" for r in restants})
    for lot in _chunks(restants, LOT_VOLUME):
        charge = [{
            "i": i,
            "article": r.item_name or r.item_code or "",
            "code": r.item_code or "",
            "groupe": r.item_group or "",
            "unite": r.uom or "",
            "description": _strip_html(r.description)[:400],
            "reperes": reperes.get(r.item_group or "", []),
        } for i, r in enumerate(lot)]
        rep = _chat_json(SYS_VOLUME, json.dumps({"lignes": charge}, ensure_ascii=False))
        vus = set()
        for item in (rep.get("lignes") or []):
            try:
                r = lot[int(item.get("i"))]
            except (TypeError, ValueError, IndexError):
                continue
            garde = normaliser_estimation(item.get("volume_m3"), item.get("pcs_carton"),
                                          item.get("volume_carton_m3"))
            if not garde:
                continue
            vol, pcs, vct = garde
            vus.add(r.name)
            propositions.append({
                "row": r.name, "idx": r.idx, "item_code": r.item_code or "",
                "item_name": r.item_name or "", "qty": qty_retenue(r),
                "uom": r.uom or "", "volume_m3": vol,
                "pcs_carton": pcs or None, "volume_carton_m3": vct or None,
                "base": (item.get("base") or "")[:160],
                "confiance": min(1, max(0, flt(item.get("confiance")))),
                "estime": 1,
            })
        for r in lot:         # ce que l'IA a laissé de côté doit se voir
            if r.name not in vus and not any(p["row"] == r.name for p in propositions):
                propositions.append({
                    "row": r.name, "idx": r.idx, "item_code": r.item_code or "",
                    "item_name": r.item_name or "", "qty": qty_retenue(r),
                    "uom": r.uom or "", "volume_m3": 0,
                    "pcs_carton": None, "volume_carton_m3": None,
                    "base": _("non estimé"), "confiance": 0, "estime": 1,
                })

    propositions.sort(key=lambda p: cint(p["idx"]))
    return {"lignes": propositions, "nb_cibles": len(cibles)}


@frappe.whitelist()
def appliquer_volumes(docname, lignes):
    """Écrit les volumes validés dans l'aperçu.

    Un volume issu de l'IA est marqué `volume_estime` : il sert au chargement
    mais ne remonte pas sur la fiche Article (cf. le controller).
    """
    _guard()
    if isinstance(lignes, str):
        lignes = json.loads(lignes)
    doc = frappe.get_doc(DOCTYPE, docname)
    par_nom = {r.name: r for r in doc.articles}

    maj = 0
    for entree in lignes or []:
        row = par_nom.get(entree.get("row"))
        vol = flt(entree.get("volume_m3"))
        if not row or vol <= 0 or vol > MAX_VOLUME_M3:
            continue
        pcs = flt(entree.get("pcs_carton"))
        vct = flt(entree.get("volume_carton_m3"))
        if pcs > 0 and vct > 0:
            row.qty_par_carton = pcs
            row.volume_carton_m3 = vct     # le controller en déduit l'unitaire
        row.volume_unitaire_m3 = vol
        row.volume_estime = cint(entree.get("estime", 1))
        maj += 1

    doc.save(ignore_permissions=True)
    frappe.db.commit()
    return {"maj": maj, "volume_total_m3": flt(doc.volume_total_m3)}
