"""Tests du renvoi au fournisseur de son propre classeur, annoté.

Convention de l'app : `unittest.TestCase` pur, données injectées, aucun accès
réseau ni base. Seules les règles de remplissage sont testées ici — l'écriture
du classeur elle-même est vérifiée à la main sur la facture Ang Ran.
"""
from __future__ import annotations

import json
import unittest

from customization_app import lci_export_fournisseur as E


class Row:
    def __init__(self, **kw):
        self.name = kw.get("name", "r")
        self.decision = ""
        self.prix_cible = 0
        self.prix_cible_negocie = 0
        self.prix_fournisseur = 0
        self.repartition_conteneurs = ""
        self.__dict__.update(kw)

    def get(self, k, default=None):
        return self.__dict__.get(k, default)


class Doc:
    def __init__(self, langue):
        self.langue_cible = langue


class TestPrixEnvoye(unittest.TestCase):
    def test_la_contre_proposition_passe_avant_tout(self):
        self.assertEqual(E._prix_cible_export(
            Row(prix_cible=33.3, prix_cible_negocie=34.0, prix_fournisseur=36.3)), 34.0)

    def test_sans_contre_proposition_on_renvoie_notre_cible_initiale(self):
        """Renvoyer SON prix dans une colonne « prix cible » contredirait
        l'observation qui, juste à côté, le dit trop cher."""
        self.assertEqual(E._prix_cible_export(
            Row(prix_cible=33.3, prix_fournisseur=36.3)), 33.3)

    def test_sans_cible_la_colonne_reste_vide(self):
        self.assertEqual(E._prix_cible_export(Row(prix_fournisseur=36.3)), 0.0)


class TestConteneurs(unittest.TestCase):
    def test_une_ligne_dans_un_seul_conteneur(self):
        self.assertEqual(E._conteneurs_ligne(
            Row(repartition_conteneurs=json.dumps([{"no": 2, "qty": 500}]))), "C2")

    def test_une_ligne_scindee_montre_ses_quantites(self):
        self.assertEqual(E._conteneurs_ligne(Row(repartition_conteneurs=json.dumps(
            [{"no": 1, "qty": 310}, {"no": 2, "qty": 590}]))), "C1 (310) + C2 (590)")

    def test_pas_de_plan_pas_de_colonne(self):
        self.assertEqual(E._conteneurs_ligne(Row()), "")

    def test_un_json_abime_ne_fait_pas_echouer_l_export(self):
        self.assertEqual(E._conteneurs_ligne(Row(repartition_conteneurs="{oups")), "")


class TestLangue(unittest.TestCase):
    def test_les_decisions_partent_dans_la_langue_du_fournisseur(self):
        self.assertEqual(E._decision(Doc("English"), "Abandonné"), "Dropped")
        self.assertEqual(E._decision(Doc("Deutsch"), "Accepté"), "Angenommen")

    def test_en_francais_la_decision_est_rendue_telle_quelle(self):
        self.assertEqual(E._decision(Doc("Français"), "À négocier"), "À négocier")

    def test_langue_inconnue_repli_francais_pour_les_entetes(self):
        self.assertEqual(E._labels(Doc("Klingon"))["qty"], "Qté retenue")

    def test_chaque_langue_declare_toutes_les_etiquettes(self):
        attendu = set(E.LABELS["Français"])
        for langue, mots in E.LABELS.items():
            self.assertEqual(set(mots), attendu, langue)
        for cle in ("photo", "retirees", "additionnels", "note_xls"):
            self.assertIn(cle, attendu)


class TestOrdreExport(unittest.TestCase):
    def test_ordre_du_document_avec_groupes_sans_abandonnees(self):
        articles = [Row(name="a", item_group="RO"), Row(name="b", item_group="RO"),
                    Row(name="c", item_group="Filtres", decision="Abandonné"),
                    Row(name="d", item_group="Filtres"), Row(name="e", item_group="")]
        ordre = E.ordre_export(articles)
        self.assertEqual([(g, (v if g == "groupe" else v.name)) for g, v in ordre],
                         [("groupe", "RO"), ("ligne", "a"), ("ligne", "b"), ("groupe", "Filtres"), ("ligne", "d"), ("ligne", "e")])

    def test_colonne_photo_la_plus_fournie(self):
        self.assertEqual(E.colonne_frequente([3, 2, 3, 3, 4, 2]), 3)
        self.assertEqual(E.colonne_frequente([2, 3]), 2)      # égalité : la plus à gauche
        self.assertIsNone(E.colonne_frequente([]))
        self.assertEqual(E.colonne_frequente([], 7), 7)

    def test_hauteur_type_mediane(self):
        self.assertEqual(E.hauteur_type([357, 20, 400, None, 380]), 380)
        self.assertEqual(E.hauteur_type([]), E.HAUTEUR_LIGNE_PHOTO)


class TestLigneModifiee(unittest.TestCase):
    def test_quantite_differente(self):
        self.assertTrue(E.est_modifiee(120, 150, 0, 36.3))

    def test_prix_cible_different(self):
        self.assertTrue(E.est_modifiee(150, 150, 33.3, 36.3))

    def test_identique_non_modifiee(self):
        self.assertFalse(E.est_modifiee(150, 150, 0, 36.3))
        self.assertFalse(E.est_modifiee(150, 150, 36.3, 36.3))

    def test_quantite_non_tranchee_ne_compte_pas(self):
        self.assertFalse(E.est_modifiee(0, 150, 0, 36.3))

    def test_additionnels_rendent_la_ligne_modifiee(self):
        self.assertTrue(E.est_modifiee(150, 150, 0, 36.3, [{"item_code": "M"}]))


class TestObservation(unittest.TestCase):
    L = E.LABELS["English"]

    def test_additionnels_decrits_avec_la_quantite_totale(self):
        adds = [{"item_name": "Membrane 1812-80 GPD", "brand": "Vontron", "qty_par_pack": 1},
                {"item_name_traduit": "Luxury faucet", "qty_par_pack": 2}]
        self.assertEqual(E.texte_additionnels(adds, 150, self.L),
                         "Additional items: + 150 × Membrane 1812-80 GPD (Vontron) ; + 300 × Luxury faucet")

    def test_observation_et_additionnels_empiles(self):
        t = E.texte_observation("Too expensive", [{"item_name": "Tap", "qty_par_pack": 1}], 10, self.L)
        self.assertEqual(t, "Too expensive\nAdditional items: + 10 × Tap")

    def test_sans_additionnels_l_observation_reste_telle_quelle(self):
        self.assertEqual(E.texte_observation("  ok ", [], 10, self.L), "ok")
        self.assertEqual(E.texte_additionnels([], 10, self.L), "")


if __name__ == "__main__":
    unittest.main()


class TestDerniereColonneUtile(unittest.TestCase):
    """Le fichier du fournisseur traîne souvent des colonnes mises en forme mais
    vides : nos colonnes doivent venir juste après son contenu, pas après ses bordures."""

    def _ws(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        ws = Workbook().active
        ws["A1"] = "titre très large"
        ws["A3"], ws["B3"], ws["J3"] = "Code", "Désignation", "Weight"
        ws["A4"], ws["J4"] = "SP-M", 150
        for c in range(11, 23):                     # K..V : fond seulement, aucune valeur
            ws.cell(row=4, column=c).fill = PatternFill("solid", fgColor="FFE0B2")
        ws.cell(row=4, column=23).value = "   "     # W : espaces = vide
        return ws

    def test_les_colonnes_seulement_formatees_ne_comptent_pas(self):
        ws = self._ws()
        self.assertGreaterEqual(ws.max_column, 23)
        self.assertEqual(E.derniere_colonne_utile(ws, premiere_ligne=3), 10)

    def test_une_valeur_isolee_plus_loin_est_gardee(self):
        ws = self._ws()
        ws.cell(row=9, column=15).value = "note"
        self.assertEqual(E.derniere_colonne_utile(ws, premiere_ligne=3), 15)

    def test_feuille_vide_repli_sur_le_defaut(self):
        from openpyxl import Workbook
        self.assertEqual(E.derniere_colonne_utile(Workbook().active, 1), 1)
